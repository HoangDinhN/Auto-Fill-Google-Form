import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_desired_distributions import (
    DOCX_PATH as FILLED_PERCENTAGES_SOURCE,
    build_distributions,
    extract_percentages,
    flatten_questions,
)


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "form-questions.json"
OUTPUT = ROOT / "Google_Form_Questions_and_Answer_Choices_Current.docx"
FALLBACK_OUTPUT = ROOT / "Google_Form_Questions_and_Answer_Choices_Updated.docx"

FONT = "Aptos"
BLACK = "000000"
MUTED = "666666"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "F3F7FB"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name=FONT, size=None, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(21)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(8)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def add_section_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_with_next = True
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_BLUE)
    p_pr.append(shd)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "120")
    ind.set(qn("w:right"), "120")
    p_pr.append(ind)
    run = paragraph.add_run(text)
    set_run_font(run, size=13.5, bold=True)
    return paragraph


def format_percentage(value):
    return f"{value:g}"


def add_question(
    doc, number, question, options=None, required=True, percentages=None
):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(f"{number}. {question}")
    set_run_font(run, size=11, bold=True)
    if required:
        marker = paragraph.add_run("  Required")
        set_run_font(marker, size=8.5, italic=True, color=MUTED)

    option_paragraphs = []
    for option_index, option in enumerate(options or []):
        option_paragraph = doc.add_paragraph()
        option_paragraph.paragraph_format.left_indent = Inches(0.24)
        option_paragraph.paragraph_format.first_line_indent = Inches(-0.02)
        option_paragraph.paragraph_format.space_after = Pt(2)
        option_paragraph.paragraph_format.tab_stops.add_tab_stop(
            Inches(6.48), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES
        )
        circle = option_paragraph.add_run("○  ")
        set_run_font(circle, name="Segoe UI Symbol", size=10.5)
        text = option_paragraph.add_run(option)
        set_run_font(text, size=10.5)
        percentage_text = (
            format_percentage(percentages[option_index])
            if percentages is not None
            else "________"
        )
        percentage = option_paragraph.add_run(f"\t{percentage_text} %")
        set_run_font(percentage, size=10.5, bold=True)
        option_paragraphs.append(option_paragraph)

    if option_paragraphs:
        paragraph.paragraph_format.keep_with_next = True
        for option_paragraph in option_paragraphs[:-1]:
            option_paragraph.paragraph_format.keep_with_next = True


def add_percentage_scale(doc, options, percentages=None):
    table = doc.add_table(rows=1, cols=len(options))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    prevent_row_split(table.rows[0])
    usable_width = 6.6
    for index, option in enumerate(options):
        cell = table.rows[0].cells[index]
        cell.width = Inches(usable_width / len(options))
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, PALE_BLUE)
        set_cell_margins(cell, top=100, start=70, bottom=100, end=70)
        label_paragraph = cell.paragraphs[0]
        label_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_paragraph.paragraph_format.space_after = Pt(3)
        label_run = label_paragraph.add_run(option)
        set_run_font(label_run, size=8.8, bold=True)

        percentage_paragraph = cell.add_paragraph()
        percentage_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        percentage_paragraph.paragraph_format.space_after = Pt(0)
        percentage_text = (
            format_percentage(percentages[index])
            if percentages is not None
            else "________"
        )
        percentage_run = percentage_paragraph.add_run(f"{percentage_text} %")
        set_run_font(percentage_run, size=10.2, bold=True)
    return table


def build_document(data):
    distributions = build_distributions(
        flatten_questions(data), extract_percentages(FILLED_PERCENTAGES_SOURCE)
    )
    percentages_by_question = {
        item["question"]: [option["percent"] for option in item["options"]]
        for item in distributions
    }

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Google Form Questions and Answer Choices")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    subtitle_run = subtitle.add_run(data["formTitle"])
    set_run_font(subtitle_run, size=11.5, italic=True)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(5)
    intro_run = intro.add_run(
        f"This document contains all {data['questionCount']} required questions and answer choices from the source Google Form. Enter a percentage beside every answer choice; the percentages for each question should total 100%."
    )
    set_run_font(intro_run, size=10.5)

    source = doc.add_paragraph()
    source.paragraph_format.space_after = Pt(12)
    label = source.add_run("Source: ")
    set_run_font(label, size=9, bold=True, color=MUTED)
    value = source.add_run(data["source"])
    set_run_font(value, size=9, color=MUTED)

    number = 1
    for section_data in data["sections"]:
        add_section_heading(doc, section_data["title"])
        questions = section_data["questions"]

        if section_data.get("type") == "multiple_choice_grid":
            for question in questions:
                percentages = percentages_by_question[question["question"]]
                add_question(
                    doc, number, question["question"], required=question["required"]
                )
                add_percentage_scale(
                    doc, section_data["sharedOptions"], percentages=percentages
                )
                number += 1
            continue

        if section_data["title"].startswith("SECTION D"):
            first = questions[0]
            add_question(
                doc,
                number,
                first["question"],
                first.get("options"),
                first["required"],
                percentages_by_question[first["question"]],
            )
            number += 1
            shared = questions[1].get("options", [])
            for question in questions[1:]:
                percentages = percentages_by_question[question["question"]]
                add_question(
                    doc, number, question["question"], required=question["required"]
                )
                add_percentage_scale(doc, shared, percentages=percentages)
                number += 1
            continue

        for question in questions:
            add_question(
                doc,
                number,
                question["question"],
                question.get("options"),
                question["required"],
                percentages_by_question[question["question"]],
            )
            number += 1

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    core = doc.core_properties
    core.title = "Google Form Questions and Answer Choices"
    core.subject = "Question and answer choice reference"
    core.author = ""
    core.keywords = "Google Forms, questionnaire, answer choices"
    try:
        doc.save(OUTPUT)
        return OUTPUT
    except PermissionError:
        doc.save(FALLBACK_OUTPUT)
        return FALLBACK_OUTPUT


if __name__ == "__main__":
    with SOURCE.open("r", encoding="utf-8") as stream:
        form_data = json.load(stream)
    output_path = build_document(form_data)
    print(output_path.name)
